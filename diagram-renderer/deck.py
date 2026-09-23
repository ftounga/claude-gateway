#!/usr/bin/env python3
"""
La construction d'une présentation .pptx (F-129 / SF-129-05).

Comme pour les schémas cloud, ce programme ne reçoit JAMAIS de code : il lit une DESCRIPTION de deck
sur son entrée standard et construit le fichier lui-même. La raison est la même — le Python d'un
modèle n'a pas à s'exécuter sur notre infrastructure.

Entrée : {"title": "...", "slides": [ {...}, ... ], "output": "/tmp/deck"}
Sortie : le chemin du .pptx écrit, sur la sortie standard.

Types de slides : "title" (titre + sous-titre), "bullets" (titre + puces), "image" (titre + image),
"table" (titre + tableau), "text" (titre + paragraphe). Chacun accepte "notes".
"""
import base64
import json
import os
import sys

MAX_SLIDES = 60
MAX_LINES = 40
MAX_LINE_CHARS = 400
MAX_IMAGES = 20
MAX_TABLE_ROWS = 20
MAX_TABLE_COLS = 8


class Refused(Exception):
    """Un refus DIT : le message explique quoi corriger."""


def text_of(raw, what, limit=MAX_LINE_CHARS):
    value = ("" if raw is None else str(raw)).strip()
    if len(value) > limit:
        raise Refused(f"{what} dépasse {limit} caractères.")
    return value


def build(spec, output):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    slides = spec.get("slides") or []
    if not slides:
        raise Refused("Aucune slide : un deck vide n'apprend rien.")
    if len(slides) > MAX_SLIDES:
        raise Refused(f"Trop de slides ({len(slides)}, maximum {MAX_SLIDES}).")

    images = spec.get("images") or {}
    if len(images) > MAX_IMAGES:
        raise Refused(f"Trop d'images ({len(images)}, maximum {MAX_IMAGES}).")

    # Les images arrivent encodées : la gateway les a lues DANS LE PROJET, sous l'isolation du tour.
    decoded = {}
    for name, payload in images.items():
        try:
            decoded[name] = base64.b64decode(payload)
        except Exception as error:  # noqa: BLE001
            raise Refused(f"Image « {name} » illisible : {error}") from error

    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]
    title_content = prs.slide_layouts[1]
    title_slide = prs.slide_layouts[0]

    for index, slide in enumerate(slides, start=1):
        kind = str(slide.get("type") or "bullets").strip().lower()
        heading = text_of(slide.get("title"), f"Le titre de la slide {index}")

        if kind == "title":
            made = prs.slides.add_slide(title_slide)
            made.shapes.title.text = heading
            subtitle = text_of(slide.get("subtitle"), f"Le sous-titre de la slide {index}")
            if subtitle and len(made.placeholders) > 1:
                made.placeholders[1].text = subtitle

        elif kind in {"bullets", "text"}:
            made = prs.slides.add_slide(title_content)
            made.shapes.title.text = heading
            lines = slide.get("bullets") if kind == "bullets" else [slide.get("text")]
            lines = [line for line in (lines or []) if line is not None]
            if len(lines) > MAX_LINES:
                raise Refused(f"Slide {index} : trop de lignes ({len(lines)}, maximum {MAX_LINES}).")
            body = made.placeholders[1].text_frame
            body.clear()
            for position, line in enumerate(lines):
                content = text_of(line, f"Une ligne de la slide {index}")
                paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
                paragraph.text = content
                paragraph.level = min(int(slide.get("level", 0) or 0), 4) if kind == "bullets" else 0

        elif kind == "image":
            made = prs.slides.add_slide(title_only)
            made.shapes.title.text = heading
            name = str(slide.get("image") or "")
            if name not in decoded:
                raise Refused(f"Slide {index} : image « {name} » absente du projet. "
                              "Un deck avec une image manquante est pire qu'un deck sans image.")
            path = f"/tmp/deck-image-{index}"
            with open(path, "wb") as handle:
                handle.write(decoded[name])
            made.shapes.add_picture(path, Inches(0.8), Inches(1.7), width=Inches(8.4))
            caption = text_of(slide.get("caption"), f"La légende de la slide {index}")
            if caption:
                box = made.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(8.4), Inches(0.6))
                box.text_frame.text = caption
                box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

        elif kind == "table":
            made = prs.slides.add_slide(title_only)
            made.shapes.title.text = heading
            rows = slide.get("rows") or []
            if not rows:
                raise Refused(f"Slide {index} : un tableau sans ligne.")
            if len(rows) > MAX_TABLE_ROWS or max(len(r or []) for r in rows) > MAX_TABLE_COLS:
                raise Refused(f"Slide {index} : tableau trop grand "
                              f"(maximum {MAX_TABLE_ROWS} lignes x {MAX_TABLE_COLS} colonnes).")
            columns = max(len(r or []) for r in rows)
            shape = made.shapes.add_table(len(rows), columns, Inches(0.6), Inches(1.7),
                                          Inches(8.8), Inches(0.8 * len(rows)))
            for r, row in enumerate(rows):
                for c in range(columns):
                    cell_value = (row or [])[c] if c < len(row or []) else ""
                    shape.table.cell(r, c).text = text_of(cell_value, f"Une cellule de la slide {index}")

        else:
            raise Refused(f"Slide {index} : type inconnu « {kind} ». "
                          "Types connus : title, bullets, text, image, table.")

        notes = text_of(slide.get("notes"), f"Les notes de la slide {index}", 4000)
        if notes:
            made.notes_slide.notes_text_frame.text = notes

    target = output if output.endswith(".pptx") else output + ".pptx"
    prs.save(target)
    return target


def main():
    try:
        spec = json.loads(sys.stdin.read() or "{}")
        if not isinstance(spec, dict):
            raise Refused("La description doit être un objet.")
        print(build(spec, spec.get("output") or "/tmp/deck"))
        return 0
    except Refused as refused:
        sys.stderr.write(str(refused))
        return 2
    except Exception as error:  # noqa: BLE001
        sys.stderr.write(f"{type(error).__name__}: {error}")
        return 3


if __name__ == "__main__":
    sys.exit(main())
