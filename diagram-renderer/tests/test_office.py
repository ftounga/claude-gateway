#!/usr/bin/env python3
"""
F-129 / SF-129-07 — le .docx et le .xlsx, vérifiés sur de VRAIS fichiers.

Ces tests construisent les fichiers et les relisent : un document qui ne s'ouvre pas ne vaut rien.
Ils tournent sans réseau ; ils ont besoin de `python-docx` et `openpyxl` (présents dans l'image du
service).

Usage : python3 tests/test_office.py
"""
import base64
import os
import sys
import tempfile
import unittest
import zlib

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import office  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def png():
    """Un PNG 1x1 valide, écrit à la main : aucun fichier de test à traîner."""
    def chunk(kind, payload):
        data = kind + payload
        return (len(payload).to_bytes(4, "big") + data
                + zlib.crc32(data).to_bytes(4, "big"))
    header = chunk(b"IHDR", (1).to_bytes(4, "big") + (1).to_bytes(4, "big")
                   + bytes([8, 2, 0, 0, 0]))
    body = chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff"))
    return b"\x89PNG\r\n\x1a\n" + header + body + chunk(b"IEND", b"")


def produced(spec):
    directory = tempfile.mkdtemp(prefix="cg-office-test-")
    return office.build(spec, os.path.join(directory, "sortie"))


def doc_spec(**extra):
    base = {
        "format": "docx",
        "title": "Compte rendu",
        "subtitle": "Revue du 26 septembre",
        "blocks": [
            {"type": "heading", "text": "Ce qui change", "level": 1},
            {"type": "text", "lines": ["Un paragraphe.", "Un autre."]},
            {"type": "bullets", "bullets": ["Un", "Deux"]},
            {"type": "table", "rows": [["Poste", "Montant"], ["S3", "12"]]},
            {"type": "pagebreak"},
            {"type": "image", "image": "archi.png", "caption": "La cible"},
        ],
        "images": {"archi.png": base64.b64encode(png()).decode("ascii")},
    }
    base.update(extra)
    return base


def sheet_spec(**extra):
    base = {
        "format": "xlsx",
        "title": "Coûts",
        "sheets": [
            {"name": "Coûts", "columns": ["Poste", "Montant"],
             "rows": [["S3", 12.5], ["EC2", 340]]},
            {"name": "Notes", "columns": ["Sujet"], "rows": [["à revoir"]]},
        ],
    }
    base.update(extra)
    return base


class DocumentTest(unittest.TestCase):

    def test_le_document_est_ouvrable_et_porte_tous_ses_blocs(self):
        """CA1 + CA3 : le .docx se relit, et chaque type de bloc a produit ce qu'il annonce."""
        path = produced(doc_spec())
        self.assertTrue(path.endswith(".docx"))
        document = Document(path)
        texts = [p.text for p in document.paragraphs]
        self.assertIn("Compte rendu", texts)
        self.assertIn("Revue du 26 septembre", texts)
        self.assertIn("Ce qui change", texts)
        self.assertIn("Un paragraphe.", texts)
        self.assertIn("La cible", texts)
        bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
        self.assertEqual(["Un", "Deux"], [p.text for p in bullets])
        self.assertEqual(1, len(document.tables))
        self.assertEqual("Montant", document.tables[0].cell(0, 1).text)
        self.assertEqual(1, len(document.inline_shapes), "l'image doit être embarquée")

    def test_les_puces_ordonnees_sont_numerotees(self):
        """CA3 : « ordered » change la liste, pas le texte."""
        spec = doc_spec(blocks=[{"type": "bullets", "bullets": ["Un"], "ordered": True}])
        document = Document(produced(spec))
        styles = [p.style.name for p in document.paragraphs if p.text == "Un"]
        self.assertEqual(["List Number"], styles)

    def test_la_charte_est_le_defaut_et_plain_ne_la_pose_pas(self):
        """CA10 : sans « theme », le titre est navy et le filet orange est posé ; « plain » non."""
        charte = Document(produced(doc_spec()))
        title = [p for p in charte.paragraphs if p.text == "Compte rendu"][0]
        self.assertEqual(office.NAVY, str(title.runs[0].font.color.rgb))
        self.assertIn(office.ORANGE, title._p.xml)

        plain = Document(produced(doc_spec(theme="plain")))
        plain_title = [p for p in plain.paragraphs if p.text == "Compte rendu"][0]
        self.assertIsNone(plain_title.runs[0].font.color.rgb)
        self.assertNotIn(office.ORANGE, plain_title._p.xml)

    def test_une_image_absente_est_refusee(self):
        """CA4 : un document avec une image manquante est pire qu'un document sans image."""
        spec = doc_spec(blocks=[{"type": "image", "image": "absente.png"}], images={})
        with self.assertRaises(office.Refused) as refused:
            produced(spec)
        self.assertIn("absente du projet", str(refused.exception))

    def test_un_type_de_bloc_inconnu_est_refuse_et_nomme_les_types_connus(self):
        """CA7 : le refus explique quoi corriger."""
        with self.assertRaises(office.Refused) as refused:
            produced(doc_spec(blocks=[{"type": "video"}]))
        message = str(refused.exception)
        self.assertIn("type inconnu", message)
        self.assertIn("bullets", message)

    def test_les_bornes_du_document_sont_tenues(self):
        """CA7 : blocs, lignes, tableau, images — chaque borne a son refus."""
        cases = {
            "Trop de blocs": doc_spec(blocks=[{"type": "pagebreak"}] * (office.MAX_BLOCKS + 1)),
            "trop de lignes": doc_spec(blocks=[{"type": "text",
                                                "lines": ["x"] * (office.MAX_LINES + 1)}]),
            "tableau trop long": doc_spec(blocks=[{"type": "table",
                                                   "rows": [["x"]] * (office.MAX_TABLE_ROWS + 1)}]),
            "tableau trop large": doc_spec(blocks=[{"type": "table",
                                                    "rows": [["x"] * (office.MAX_TABLE_COLS + 1)]}]),
            "Trop d'images": doc_spec(images={"i%d" % i: "" for i in range(office.MAX_IMAGES + 1)}),
            "dépasse": doc_spec(blocks=[{"type": "text",
                                         "lines": ["x" * (office.MAX_LINE_CHARS + 1)]}]),
        }
        for expected, spec in cases.items():
            with self.subTest(expected):
                with self.assertRaises(office.Refused) as refused:
                    produced(spec)
                self.assertIn(expected, str(refused.exception))

    def test_un_document_sans_bloc_est_refuse(self):
        with self.assertRaises(office.Refused):
            produced(doc_spec(blocks=[]))


class SpreadsheetTest(unittest.TestCase):

    def test_le_classeur_est_ouvrable_et_les_nombres_restent_des_nombres(self):
        """CA2 : feuilles nommées, en-tête figé et filtrable, montants sommables."""
        path = produced(sheet_spec())
        self.assertTrue(path.endswith(".xlsx"))
        workbook = load_workbook(path)
        self.assertEqual(["Coûts", "Notes"], workbook.sheetnames)
        sheet = workbook["Coûts"]
        self.assertEqual("Poste", sheet["A1"].value)
        self.assertEqual(12.5, sheet["B2"].value)
        self.assertIsInstance(sheet["B3"].value, int)
        self.assertEqual("A2", sheet.freeze_panes)
        self.assertEqual("A1:B3", sheet.auto_filter.ref)
        self.assertGreater(sheet.column_dimensions["A"].width, 0)

    def test_la_charte_habille_l_entete_et_plain_ne_l_habille_pas(self):
        """CA10 : l'en-tête navy est le défaut ; « plain » rend le gabarit neutre."""
        charte = load_workbook(produced(sheet_spec()))["Coûts"]
        self.assertIn(office.NAVY, str(charte["A1"].fill.fgColor.rgb))
        self.assertTrue(charte["A1"].font.bold)

        plain = load_workbook(produced(sheet_spec(theme="plain")))["Coûts"]
        self.assertNotIn(office.NAVY, str(plain["A1"].fill.fgColor.rgb))
        self.assertFalse(plain["A1"].font.bold)

    def test_le_nom_d_onglet_est_nettoye_tronque_et_desambiguise(self):
        """Un classeur refusé à l'ouverture vaut moins qu'un onglet renommé."""
        spec = sheet_spec(sheets=[
            {"name": "Coûts/2026[T1]", "columns": ["A"], "rows": [["x"]]},
            {"name": "Coûts/2026[T1]", "columns": ["A"], "rows": [["x"]]},
            {"name": "x" * 60, "columns": ["A"], "rows": [["x"]]},
        ])
        names = load_workbook(produced(spec)).sheetnames
        self.assertEqual("Coûts 2026 T1", names[0])
        self.assertEqual("Coûts 2026 T1 (2)", names[1])
        self.assertEqual(office.MAX_SHEET_NAME, len(names[2]))

    def test_les_bornes_du_classeur_sont_tenues(self):
        cases = {
            "Trop de feuilles": sheet_spec(sheets=[{"columns": ["A"], "rows": [["x"]]}]
                                           * (office.MAX_SHEETS + 1)),
            "trop de lignes": sheet_spec(sheets=[{"columns": ["A"],
                                                  "rows": [["x"]] * (office.MAX_ROWS + 1)}]),
            "trop de colonnes": sheet_spec(sheets=[{"columns": ["A"] * (office.MAX_COLS + 1),
                                                    "rows": []}]),
            "dépasse": sheet_spec(sheets=[{"columns": ["A"],
                                           "rows": [["x" * (office.MAX_CELL_CHARS + 1)]]}]),
            "ni colonnes ni lignes": sheet_spec(sheets=[{"name": "vide"}]),
        }
        for expected, spec in cases.items():
            with self.subTest(expected):
                with self.assertRaises(office.Refused) as refused:
                    produced(spec)
                self.assertIn(expected, str(refused.exception))

    def test_un_classeur_sans_feuille_est_refuse(self):
        with self.assertRaises(office.Refused):
            produced(sheet_spec(sheets=[]))


class FormatTest(unittest.TestCase):

    def test_un_format_inconnu_est_refuse_et_nomme_les_formats_connus(self):
        with self.assertRaises(office.Refused) as refused:
            produced({"format": "pdf", "blocks": [{"type": "pagebreak"}]})
        self.assertIn("docx", str(refused.exception))

    def test_un_theme_inconnu_est_refuse(self):
        with self.assertRaises(office.Refused) as refused:
            produced(doc_spec(theme="fluo"))
        self.assertIn("Thème inconnu", str(refused.exception))


if __name__ == "__main__":
    unittest.main(verbosity=2)
